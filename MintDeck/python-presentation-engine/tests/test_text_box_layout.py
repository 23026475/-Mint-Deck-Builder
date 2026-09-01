from __future__ import annotations

import unittest

from pptx.util import Inches

from presentation_engine.services.text_box_layout import TextBoxLayoutPolicy


CONFIG = {
    "defaults": {
        "mode": "fixed",
        "margin_left": 0.08,
        "margin_right": 0.08,
        "margin_top": 0.05,
        "margin_bottom": 0.05,
        "line_height_multiplier": 1.2,
        "paragraph_gap_lines": 0.35,
        "characters_per_point_inch": 27.8,
    },
    "archetypes": {
        "statement": {
            "support": {
                "idx": 2,
                "mode": "dynamic",
                "min_height": 0.4,
                "max_height": 1.65,
            }
        },
        "image_right": {
            "body": {
                "idx": 12,
                "mode": "dynamic",
                "min_height": 0.55,
                "max_height": 1.5,
            }
        },
        "process_flow": {
            "step1_body": {
                "idx": 21,
                "mode": "fixed",
            }
        },
    },
}


class FakeTextFrame:
    def __init__(self) -> None:
        self.paragraphs = []
        self.vertical_anchor = None


class FakeShape:
    """Explicit-geometry shape used only for isolated policy unit tests."""

    def __init__(
        self,
        width: float,
        height: float,
        *,
        left: float = 0.79,
        top: float = 1.53,
    ) -> None:
        self.name = "Fake Text Placeholder"
        self.has_text_frame = True
        self.text_frame = FakeTextFrame()
        self.left = Inches(left)
        self.top = Inches(top)
        self.width = Inches(width)
        self.height = Inches(height)


class TextBoxLayoutPolicyTests(unittest.TestCase):
    def setUp(self) -> None:
        self.policy = TextBoxLayoutPolicy(config_data=CONFIG)

    def test_statement_support_grows_for_multiple_paragraphs(self):
        shape = FakeShape(9.0, 0.4, left=0.9, top=4.4)
        original_left = shape.left
        original_top = shape.top
        original_width = shape.width

        result = self.policy.apply(
            shape,
            archetype="statement",
            placeholder_name="support",
            text=(
                "IT capability is expanding beyond traditional infrastructure into software, data, analytics and AI.\n"
                "Technology teams now need strong foundations, practical engineering discipline, reliable data practices.\n"
                "Takeaway: build a balanced skills portfolio."
            ),
            font_size_pt=15.0,
        )

        self.assertIsNotNone(result)
        self.assertEqual(shape.left, original_left)
        self.assertEqual(shape.top, original_top)
        self.assertEqual(shape.width, original_width)
        self.assertGreater(result.applied_height_inches, 0.4)
        self.assertLessEqual(result.applied_height_inches, 1.65)

    def test_image_body_shrinks_to_configured_minimum(self):
        shape = FakeShape(7.01, 5.14)
        original_left = shape.left
        original_top = shape.top
        original_width = shape.width

        result = self.policy.apply(
            shape,
            archetype="image_right",
            placeholder_name="body",
            text="Capability grows through shared learning and practical collaboration.",
            font_size_pt=14.0,
        )

        self.assertIsNotNone(result)
        self.assertEqual(shape.left, original_left)
        self.assertEqual(shape.top, original_top)
        self.assertEqual(shape.width, original_width)
        self.assertAlmostEqual(result.applied_height_inches, 0.55, places=2)

    def test_fixed_process_box_is_not_resized(self):
        shape = FakeShape(2.13, 1.6)
        original_geometry = (
            shape.left,
            shape.top,
            shape.width,
            shape.height,
        )

        result = self.policy.apply(
            shape,
            archetype="process_flow",
            placeholder_name="step1_body",
            text="Reliable platforms and secure access",
            font_size_pt=12.5,
        )

        self.assertIsNone(result)
        self.assertEqual(
            (shape.left, shape.top, shape.width, shape.height),
            original_geometry,
        )

    def test_unknown_placeholder_is_not_resized(self):
        shape = FakeShape(5.0, 1.0)
        original_geometry = (
            shape.left,
            shape.top,
            shape.width,
            shape.height,
        )

        result = self.policy.apply(
            shape,
            archetype="cover_dark",
            placeholder_name="unknown",
            text="Text",
            font_size_pt=14.0,
        )

        self.assertIsNone(result)
        self.assertEqual(
            (shape.left, shape.top, shape.width, shape.height),
            original_geometry,
        )


if __name__ == "__main__":
    unittest.main()
